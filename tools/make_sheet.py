#!/usr/bin/env python3
"""Build the submission workbook from a pipeline run.

  pip install openpyxl            # only this script needs it; the pipeline itself is stdlib-only
  python3 tools/make_sheet.py out/agencies out/agencies/inboxkit_lead_list.xlsx

Tabs: "Lead list" (100 rows, brief's columns first, India rows first) and "Scale & cost"
(the 2,000-row paragraph plus a formula-driven cost model; blue cells are inputs).
"""
import csv
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BOLD = Font(name=FONT, bold=True)
NORMAL = Font(name=FONT)
HEAD_FILL = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF")
INPUT_FILL = PatternFill("solid", fgColor="FFF59D")
WRAP = Alignment(wrap_text=True, vertical="top")

LEAD_COLS = [
    ("Name", "name", 22), ("Title", "title", 28), ("Company", "company", 24), ("Domain", "domain", 24),
    ("LinkedIn URL", "linkedin_url", 42), ("Email", "email", 30), ("Verification status", "verification_status", 18),
    ("Signal that qualified them", "qualifying_signal", 60), ("Source", "source", 48),
    ("First line (workflow)", "first_line", 70), ("First line basis", "line_basis", 14),
    ("Signal line (Clay prompt, not yet run)", None, 34), ("Country", "country", 16),
]


def lead_tab(wb, rows):
    ws = wb.active
    ws.title = "Lead list"
    for i, (h, _, w) in enumerate(LEAD_COLS, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.alignment = HEAD_FONT, HEAD_FILL, WRAP
        ws.column_dimensions[get_column_letter(i)].width = w
    rows = sorted(rows, key=lambda r: (r["country"] != "India", r["company"].lower(), r["name"]))
    for r_i, r in enumerate(rows, 2):
        for c_i, (_, key, _) in enumerate(LEAD_COLS, 1):
            cell = ws.cell(row=r_i, column=c_i, value=r.get(key, "") if key else "")
            cell.font = NORMAL
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(LEAD_COLS)), len(rows) + 1)
    ws["G1"].comment = Comment(
        "pattern_mx_ok = guessed first@domain, domain has working MX, mailbox not verified. "
        "No paid finder or verifier was used for this list.", "workflow")
    return len(rows)


# (label, value or formula, number format, is_input, note)
MODEL = [
    ("Volume", None, None, False, None),
    ("Final contacts in the list", 2000, "#,##0", True, "Target from the brief."),
    ("Pipeline pass rate, raw rows to final", 0.725, "0.0%", True,
     "From this run: 100 final rows from 138 raw Clay rows (rejects.csv lists why)."),
    ("Prospeo find rate (verified email found)", 0.75, "0.0%", True,
     "Assumption. Replace with Prospeo's hit rate on the first 200 rows."),
    ("Contacts needing an email lookup", "=B3/B5", "#,##0", False, None),
    ("Raw Clay rows to pull", "=B6/B4", "#,##0", False,
     "Clay Launch includes unlimited search (clay.com/pricing), so pulling rows is not the cost."),
    ("Clay: Launch plan", None, None, False, None),
    ("Launch plan, per month ($)", 185, "$#,##0", True,
     "User's plan choice (quoted as $184). clay.com/pricing lists Launch at $185/mo monthly, $167/mo annual."),
    ("Data credits included per month", 2500, "#,##0", True, "clay.com/pricing: Launch, 2.5K data credits/mo."),
    ("Credits per row for the signal-line prompt", 5, "#,##0", True,
     "User estimate: 4 to 5 credits per row. Upper end used. If Claygent bills Actions instead (15K/mo on Launch), overage falls to zero."),
    ("Rows run through the signal prompt", "=B3", "#,##0", False, "Run on final contacts only, after dedupe and checks."),
    ("Credits needed", "=B11*B12", "#,##0", False, None),
    ("Credits over the plan", "=MAX(0,B13-B10)", "#,##0", False, None),
    ("Top-up price per data credit ($)", 0.065, "$0.000", True,
     "clay.com/pricing: data credits from $0.05, Launch top-ups carry a 30% premium, so $0.065."),
    ("Clay total ($)", "=B9+B14*B15", "$#,##0", False, None),
    ("Prospeo: email finding", None, None, False, None),
    ("Starter plan, per month ($)", 49, "$#,##0", True,
     "User's plan choice (base plan). 2026 third-party pricing reviews list Starter at $49/mo with 2,000 credits; confirm on prospeo.io."),
    ("Starter credits per month", 2000, "#,##0", True, "Same source as above."),
    ("Next plan (Growth), per month ($)", 99, "$#,##0", True, "Growth: $99/mo, 5,000 credits (same source)."),
    ("Credits per found email", 1, "#,##0", True, "1 credit per verified email found; assumed charged only on finds."),
    ("Credits needed", "=B3*B21", "#,##0", False, None),
    ("Prospeo total ($)", "=IF(B22<=B19,B18,B20)", "$#,##0", False, None),
    ("ZeroBounce: second-pass verification", None, None, False, None),
    ("Pay-as-you-go pack price ($)", 39, "$#,##0", True,
     "zerobounce.net pricing: pay as you go starts at $39 for 2,000 validations (minimum purchase)."),
    ("Validations per pack", 2000, "#,##0", True, "Same source."),
    ("Emails to verify", "=B3", "#,##0", False, "Every found email is re-checked before it enters a sequence."),
    ("ZeroBounce total ($)", "=ROUNDUP(B27/B26,0)*B25", "$#,##0", False, None),
    ("HeyReach: LinkedIn outreach", None, None, False, None),
    ("Growth plan, per sender per month ($)", 79, "$#,##0", True, "heyreach.io/pricing: Growth $79 per sender, monthly billing."),
    ("LinkedIn sender accounts", 3, "#,##0", True, "From the memo: 3 senders, about 70 connects a week each."),
    ("HeyReach total ($)", "=B30*B31", "$#,##0", False, None),
    ("InboxKit: sending infrastructure", None, None, False, None),
    ("Emails per contact", 4, "#,##0", True, "Sequence in copy/: 4 emails."),
    ("Sending days", 15, "#,##0", True, "3-week plan, business days."),
    ("Peak day vs average", 1.1, "0.00", True, "Daily cohorts overlap with follow-ups; about 10% above average."),
    ("Peak sends per day", "=B3*B34/B35*B36", "#,##0", False, None),
    ("Cold emails per mailbox per day", 15, "#,##0", True,
     "inboxkit.com/email-warmup: 10 to 15 per mailbox per day after warmup, max 3 mailboxes per domain."),
    ("Mailboxes", "=ROUNDUP(B37/B38,0)", "#,##0", False, None),
    ("Domains (3 mailboxes each)", "=ROUNDUP(B39/3,0)", "#,##0", False, None),
    ("Pre-warmed mailbox, per month ($)", 6, "$#,##0.00", True,
     "inboxkit.com/pricing: pre-warmed Google/Microsoft mailboxes from $6/mo (2 to 4 weeks warmed)."),
    ("Pre-warmed domain, one time ($)", 12.5, "$#,##0.00", True, "inboxkit.com/pricing: pre-warmed domain fee from $12.50."),
    ("InboxKit total ($)", "=B39*B41+B40*B42", "$#,##0", False, None),
    ("Total", None, None, False, None),
    ("Total for 2,000 contacts, first month ($)", "=B16+B23+B28+B32+B43", "$#,##0", False, None),
    ("Cost per contact, all in ($)", "=B45/B3", "$0.00", False, None),
    ("Cost per contact, data only: Clay + Prospeo + ZeroBounce ($)", "=(B16+B23+B28)/B3", "$0.00", False, None),
]

PARAGRAPH = (
    "Scaling to 2,000 rows. The pipeline does not change: pull about 3,700 raw rows from Clay search, run "
    "`run.py build` (it took under a minute and no network in --offline mode for 138 rows), and the same "
    "dedupe, suppression and email checks apply. What changes is where each email comes from. At 2,000 rows "
    "the free first@domain guess is not good enough to send on, so Prospeo finds a verified address for each "
    "contact, ZeroBounce re-checks every one before it enters a sequence, and the Clay signal prompt writes "
    "the opener for every final row. Using the plans chosen (Clay Launch, Prospeo Starter, ZeroBounce pay as "
    "you go, three HeyReach senders, pre-warmed InboxKit mailboxes) the model below comes to about $0.38 per "
    "contact for data and about $0.71 all in for the first month. The biggest lever is Clay credits: if the "
    "signal prompt bills against Actions rather than Data Credits, or runs only on rows with no free signal, "
    "the data cost roughly halves. Every blue cell is an input and the totals recalculate."
)


def cost_tab(wb):
    ws = wb.create_sheet("Scale & cost")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 90
    ws["A1"] = PARAGRAPH
    ws["A1"].font, ws["A1"].alignment = NORMAL, WRAP
    ws.merge_cells("A1:C1")
    ws.row_dimensions[1].height = 150
    row = 1
    for label, val, fmt, is_input, note in MODEL:
        row += 1
        a = ws.cell(row=row, column=1, value=label)
        if val is None:
            a.font = BOLD
            continue
        a.font = NORMAL
        b = ws.cell(row=row, column=2, value=val)
        b.font = BLUE if is_input else (BOLD if label.startswith(("Total", "Cost per")) else NORMAL)
        if is_input:
            b.fill = INPUT_FILL
        b.number_format = fmt
        if note:
            c = ws.cell(row=row, column=3, value=note)
            c.font, c.alignment = Font(name=FONT, italic=True, color="555555"), WRAP
    ws.cell(row=row + 2, column=1, value="Legend: blue on yellow = input you can change; black = formula.").font = Font(
        name=FONT, italic=True)
    return ws


def main(out_dir, path):
    rows = list(csv.DictReader(open(out_dir.rstrip("/") + "/lead_list.csv")))
    wb = Workbook()
    n = lead_tab(wb, rows)
    cost_tab(wb)
    wb.save(path)
    print("Wrote %s: %d leads" % (path, n))


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
